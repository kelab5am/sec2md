"""Optional XLSX rendering. No spreadsheet dependency is imported at module import."""
from __future__ import annotations

from dataclasses import dataclass
from importlib.metadata import version
from io import BytesIO
import math
import re
from typing import Sequence
from urllib.parse import quote, urldefrag

from sec2md.xlsx_tables import PreparedTable

MAX_ROWS = 1048576
MAX_COLUMNS = 16384
# Small visible chunks also avoid Excel's maximum row height clipping long text.
TEXT_CHUNK = 400
INVALID_XML = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]')


@dataclass(frozen=True)
class WorksheetResult:
    ordinal: int
    worksheet_name: str
    status: str
    issues: tuple[str, ...]
    copy_range: str | None = None
    original_range: str | None = None


def _sheet_name(table, used):
    prefix = f'{table.source.ordinal:03d}_p{table.source.page:03d}_'
    title = re.sub(r'[\x00-\x1f\[\]:*?/\\]', '', INVALID_XML.sub('', table.title)).strip(" '") or 'Table'
    base = (prefix + title)[:31].rstrip("'")
    name, index = base, 2
    while name.lower() in used or name.lower() in {'contents', 'history'}:
        suffix = f'_{index}'
        name = base[:31-len(suffix)] + suffix
        index += 1
    used.add(name.lower())
    return name


def render_workbook(tables: Sequence[PreparedTable], *, source_url: str | None,
                    source_hash: str, hash_kind: str,
                    document_diagnostics: Sequence[str] = ()) -> tuple[bytes, tuple[WorksheetResult, ...]]:
    """Render in supplied traversal order; report writer safeguards for strict callers."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.hyperlink import Hyperlink
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise ImportError('XLSX export requires pip install "sec2md[xlsx]"') from exc

    wb = Workbook()
    contents = wb.active
    contents.title = 'Contents'
    navy = '253D55'
    used = {'contents'}
    writer_issues = {}

    def safe_text(ws, text):
        if INVALID_XML.search(text):
            writer_issues.setdefault(ws.title, set()).add(
                'XML-incompatible characters escaped as \\uXXXX; backslashes doubled in affected text.')
            text = INVALID_XML.sub(lambda m: f'\\u{ord(m[0]):04x}', text.replace('\\', '\\\\'))
        return text

    def put(ws, row, col, value, *, heading=False, title=False, number_format='@'):
        if row >= MAX_ROWS:
            writer_issues.setdefault(ws.title, set()).add(
                'Content unavailable in full because Excel row limit was exceeded.')
            row, col = MAX_ROWS, 1
            value = 'UNAVAILABLE IN FULL: Excel row limit reached. Consult supplied source.'
            heading, title, number_format = True, False, '@'
        cell = ws.cell(row, col)
        if isinstance(value, str):
            # Force type after assignment: source formulas and error tokens stay text.
            cell.value = safe_text(ws, value)
            cell.data_type = 's'
        else:
            cell.value = value
        cell.number_format = number_format
        cell.font = Font(name='Arial', size=17 if title else 11,
                         bold=heading or title, color='FFFFFF' if heading else navy)
        cell.alignment = Alignment(vertical='top', wrap_text=True,
                                   horizontal='right' if number_format != '@' else 'left')
        if heading:
            cell.fill = PatternFill('solid', fgColor=navy)
        elif row % 2 == 0:
            cell.fill = PatternFill('solid', fgColor='F5F7F9')
        width = ws.column_dimensions[get_column_letter(col)].width or 25
        lines = sum(max(1, math.ceil(len(line) / max(1, int(width * .85))))
                    for line in str(value or '').split('\n'))
        height = min(409, max(24, lines * (23 if title else 16) + 8))
        ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 0, height)
        return cell

    def block(ws, row, text, *, heading=False, title=False):
        text = safe_text(ws, text)
        for part in (text[i:i+TEXT_CHUNK] for i in range(0, len(text), TEXT_CHUNK)):
            put(ws, row, 1, part, heading=heading, title=title)
            if row >= MAX_ROWS:
                return MAX_ROWS
            row += 1
        return row

    def link(cell, *, target=None, location=None):
        cell.hyperlink = Hyperlink(ref=cell.coordinate, target=target, location=location)
        cell.font = Font(name='Arial', size=11, color='0563C1', underline='single')

    def internal(name, coordinate):
        return f"'{name.replace(chr(39), chr(39)*2)}'!{coordinate}"

    def setup(ws, columns):
        ws.sheet_view.showGridLines = False
        ws.column_dimensions['A'].width = 65
        for col in range(2, columns + 1):
            ws.column_dimensions[get_column_letter(col)].width = 25

    setup(contents, 8)
    row = block(contents, 1, 'Filing tables', title=True)
    for label, text in [('Source', source_url or 'Supplied document'), ('SHA-256 kind', hash_kind),
                        ('SHA-256', source_hash), ('sec2md version', version('sec2md')),
                        ('Export schema', '1')]:
        put(contents, row, 1, label)
        # Arbitrary source metadata can be long as well.
        for part in (text[i:i+TEXT_CHUNK] for i in range(0, len(text), TEXT_CHUNK)):
            put(contents, row, 2, part)
            row += 1
    for diagnostic in document_diagnostics:
        row = block(contents, row, diagnostic)
    summary_row = row
    row += 2
    for col, text in enumerate(('Table / copy area', 'Originals', 'Parser page', 'Printed page',
                                'Status', 'Review count'), 1):
        put(contents, row, col, text, heading=True)
    row += 1
    records = []
    for table in tables:
        name = _sheet_name(table, used)
        ws = wb.create_sheet(name)
        ws.freeze_panes = 'B1'
        issues = list(table.issues)
        width = max(len(table.headers), max((len(r) for r in table.rows), default=0))
        original_width = max((len(r) for r in table.original_rows), default=0)
        # Conservative estimate includes originals, context, provenance and expansion.
        strings = [table.title, table.units, table.source.original_text, *table.headers,
                   *table.notes, *table.issues, *table.source.context_before,
                   *table.source.context_after, *(c.text for c in table.source.source_cells),
                   *(str(c.value or '') for r in table.rows for c in r)]
        estimate = 100 + len(table.rows) + len(table.original_rows) + len(table.source.source_cells) * 3 + sum(
            math.ceil(len(s) / TEXT_CHUNK) + 1 for s in strings)
        estimate += sum(len(r) for r in table.cell_sources)
        oversized = max(width, original_width) > MAX_COLUMNS or estimate > MAX_ROWS
        if oversized:
            issues.append('Excel dimension limit: copy grid unavailable; original text retained when space permits.')
        chunked = any(len(s) > TEXT_CHUNK for s in (*table.headers,
                      *(str(c.value or '') for body in table.rows for c in body)))
        if chunked:
            issues.append('Long text is retained in labeled chunks; concatenate without separators.')
        status = 'source_text_only' if oversized else table.status
        if issues and status == 'exported':
            status = 'needs_review'
        setup(ws, max(3, min(max(width, original_width), MAX_COLUMNS)) if not oversized else 3)
        r = block(ws, 1, table.title or 'Table', title=True)
        put(ws, r, 1, 'Back to Contents')
        link(ws.cell(r, 1), location="'Contents'!A1")
        r += 1
        if source_url:
            r = block(ws, r, source_url)
            destination = source_url
            if table.source.source_anchor:
                destination = urldefrag(source_url)[0] + '#' + quote(table.source.source_anchor, safe='')
            link(ws.cell(r-1, 1), target=destination)
        pages = f'Parser page {table.source.page}'
        if table.source.display_page is not None:
            pages += f' | Printed page {table.source.display_page}'
        r = block(ws, r, pages)
        if table.source.element_id:
            r = block(ws, r, f'Local element: {table.source.element_id}')
        status_row = r
        r = block(ws, r, f'Status: {status}')
        instruction = r
        r += 1
        copy_start = r
        r = block(ws, r, table.units or 'Units not established')
        copy_range = None
        long_copy = []
        if status != 'source_text_only':
            for col, header in enumerate(table.headers, 1):
                text = header if len(header) <= TEXT_CHUNK else f'Long header: see notes ({get_column_letter(col)}{r})'
                put(ws, r, col, text, heading=True)
                if len(header) > TEXT_CHUNK:
                    long_copy.append((f'Header {get_column_letter(col)}{r}', header))
            r += 1
            for body in table.rows:
                for col, item in enumerate(body, 1):
                    value = item.value
                    if isinstance(value, str) and len(value) > TEXT_CHUNK:
                        long_copy.append((f'Copy cell {get_column_letter(col)}{r}', value))
                        value = 'Long text: see notes / original text'
                    put(ws, r, col, value, number_format=item.number_format)
                    if item.review_reason:
                        long_copy.append((f'Review {get_column_letter(col)}{r}', item.review_reason))
                r += 1
            copy_range = f'A{copy_start}:{get_column_letter(max(1,width))}{r-1}'
            put(ws, instruction, 1, f'Copy {copy_range}, including units and complete headers.')
        else:
            put(ws, instruction, 1, 'Copy grid unavailable — source text only', heading=True)
        r += 1
        r = block(ws, r, 'Notes and review', heading=True)
        originals_link_row = r
        r = block(ws, r, 'Go to original text')
        r = block(ws, r, 'Originals retain extracted text and source spans. Parser whitespace normalization applies; consult retained HTML for exact visual appearance.')
        for note in (*issues, *table.notes):
            r = block(ws, r, note)
        for label, text in long_copy:
            r = block(ws, r, label + ' (chunks; concatenate without separators)')
            r = block(ws, r, text)
        for context in (*table.source.context_before, *table.source.context_after):
            r = block(ws, r, 'Nearby source context')
            r = block(ws, r, context)
        for label, target in table.references:
            r = block(ws, r, label)
            r = block(ws, r, target)
            if target:
                link(ws.cell(r-1, 1), target=target)
        original_start = min(r + 1, MAX_ROWS)
        link(ws.cell(min(originals_link_row, MAX_ROWS), 1), location=internal(name, f'A{original_start}'))
        if not oversized:
            r = block(ws, original_start, 'Original cell grid — source spans recorded below', heading=True)
            for original_row in table.original_rows:
                for col, text in enumerate(original_row, 1):
                    put(ws, r, col, text if len(text) <= TEXT_CHUNK else 'Long text: see source-cell chunks below')
                r += 1
            r += 1
        else:
            r = original_start
        r = block(ws, r, 'Original extracted text (chunks; concatenate without separators)', heading=True)
        text = table.source.original_text
        if r + math.ceil(len(text) / TEXT_CHUNK) + 3 > MAX_ROWS:
            r = block(ws, r, 'UNAVAILABLE IN FULL: original text exceeds Excel row limits. Consult supplied source.')
            issues.append('Original text unavailable in full because Excel row limit was exceeded.')
        else:
            r = block(ws, r, text)
        r += 1
        # A lossless source-cell ledger explicitly retains coordinates and spans.
        if not oversized:
            r = block(ws, r, 'Original source cells — zero-based row/column and spans', heading=True)
            for cell in table.source.source_cells:
                r = block(ws, r, f'row {cell.row}, column {cell.column}; rowspan {cell.rowspan}, colspan {cell.colspan}')
                r = block(ws, r, cell.text)
                for label, target in cell.links:
                    r = block(ws, r, f'Reference: {label}')
                    r = block(ws, r, target)
                    if target:
                        link(ws.cell(r-1, 1), target=target)
            r = block(ws, r, 'Copy cell source origins (zero-based row, column)', heading=True)
            for ri, source_row in enumerate(table.cell_sources):
                for ci, origins in enumerate(source_row):
                    r = block(ws, r, f'Body row {ri}, column {ci}: {origins}')
        issues.extend(sorted(writer_issues.get(name, ())))
        if writer_issues.get(name):
            status = 'source_text_only' if r >= MAX_ROWS else 'needs_review'
            for issue in sorted(writer_issues[name]):
                r = block(ws, r, issue)
        put(ws, status_row, 1, f'Status: {status}')
        original_end = get_column_letter(max(1, original_width)) if not oversized else 'A'
        original_range = f'A{original_start}:{original_end}{min(r-1, MAX_ROWS)}'
        record = WorksheetResult(table.source.ordinal, name, status, tuple(dict.fromkeys(issues)),
                                 copy_range, original_range)
        records.append(record)
        # Titles are retained in full on the table sheet, even for long Contents labels.
        put(contents, row, 1, table.title if len(table.title) <= TEXT_CHUNK else name)
        link(contents.cell(row, 1), location=internal(name, f'A{copy_start}'))
        put(contents, row, 2, 'Original text')
        link(contents.cell(row, 2), location=internal(name, f'A{original_start}'))
        put(contents, row, 3, table.source.page, number_format='0')
        put(contents, row, 4, table.source.display_page, number_format='0')
        put(contents, row, 5, status)
        put(contents, row, 6, len(record.issues), number_format='0')
        row += 1
    summary = 'No tables detected' if not records else f'{len(records)} tables: ' + ', '.join(
        f'{sum(r.status == state for r in records)} {state}'
        for state in ('exported', 'needs_review', 'source_text_only'))
    if document_diagnostics:
        summary += ' | Document diagnostics require review'
    put(contents, summary_row, 1, summary)
    output = BytesIO()
    wb.save(output)
    return output.getvalue(), tuple(records)
