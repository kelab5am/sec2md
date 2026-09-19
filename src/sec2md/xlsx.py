"""One-document XLSX export with optional dependencies and safe publication."""
from __future__ import annotations

from dataclasses import dataclass
import hashlib
import os
from pathlib import Path
import tempfile
from typing import Literal

from sec2md.core import _link_resolution_url, _resolve_source
from sec2md.parser import Parser
from sec2md.quality import build_diagnostics, enforce_quality
from sec2md.utils import is_url
from sec2md.xlsx_tables import prepare_table
from sec2md.xlsx_writer import render_workbook


@dataclass(frozen=True)
class XlsxTableResult:
    ordinal: int
    sheet_name: str
    status: Literal['exported', 'needs_review', 'source_text_only']
    issues: tuple[str, ...]


@dataclass(frozen=True)
class XlsxExportResult:
    path: Path
    status: Literal['complete', 'needs_review', 'no_tables']
    tables: tuple[XlsxTableResult, ...]
    diagnostics: tuple[str, ...]


class XlsxDependencyError(ImportError):
    """The optional XLSX dependency is unavailable."""


class XlsxQualityError(ValueError):
    """Strict export rejected table reliability issues before publication."""

    def __init__(self, issues: tuple[str, ...]):
        self.issues = issues
        super().__init__('; '.join(issues))


def _publish(payload: bytes, destination: Path, *, overwrite: bool, load_workbook) -> None:
    """Validate a sibling staging file, then link without clobber or replace."""
    temp_path = None
    try:
        with tempfile.NamedTemporaryFile(
            mode='wb', dir=destination.parent, prefix=f'.{destination.name}.',
            suffix='.xlsx', delete=False,
        ) as staged:
            temp_path = Path(staged.name)
            staged.write(payload)
        # Load all worksheets rather than merely opening the ZIP container.
        with temp_path.open('rb') as saved:
            workbook = load_workbook(saved)
            workbook.close()
        if overwrite:
            os.replace(temp_path, destination)
        else:
            # Hard-link creation fails if any destination entry already exists.
            # Unsupported filesystems fail explicitly; there is no copy fallback.
            os.link(temp_path, destination)
    finally:
        if temp_path is not None:
            temp_path.unlink(missing_ok=True)


def export_xlsx(
    source: str | bytes,
    destination: str | Path,
    *,
    base_url: str | None = None,
    user_agent: str | None = None,
    quality_policy: Literal['strict', 'warn', 'off'] = 'warn',
    overwrite: bool = False,
) -> XlsxExportResult:
    """Export one HTML document; read local files as bytes before calling.

    Raw input never fetches linked content. URL input fetches only that document.
    Strict quality errors occur before publication. No-overwrite publication
    requires filesystem hard-link support and raises OSError if unavailable.
    """
    if not isinstance(quality_policy, str) or quality_policy not in {'strict', 'warn', 'off'}:
        raise ValueError(f'invalid quality_policy: {quality_policy}')
    if not isinstance(source, (str, bytes)):
        raise TypeError('source must be an HTML string, bytes, or URL string')
    source_url = source if isinstance(source, str) and is_url(source) else None
    link_url = _link_resolution_url(source_url, base_url)
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise XlsxDependencyError('XLSX export requires pip install "sec2md[xlsx]"') from exc
    path = Path(destination)
    if path.is_dir():
        raise IsADirectoryError(f'XLSX destination is a directory: {path}')
    if not path.parent.exists():
        raise FileNotFoundError(f'XLSX destination parent does not exist: {path.parent}')
    if not path.parent.is_dir():
        raise NotADirectoryError(f'XLSX destination parent is not a directory: {path.parent}')
    if not overwrite and os.path.lexists(path):
        raise FileExistsError(f'XLSX destination already exists: {path}')

    html, decode_diagnostics = _resolve_source(source, user_agent=user_agent)
    parser = Parser(html, source_url=link_url, decode_diagnostics=decode_diagnostics,
                    capture_tables=True)
    pages = parser.get_pages(include_images=False)
    diagnostics = parser.diagnostics
    if diagnostics is None:
        diagnostics = build_diagnostics(
            html, '\n\n'.join(page.content for page in pages if page.content), pages,
            mapped_element_ids=tuple(key for key, nodes in parser.block_nodes_map.items() if nodes),
            trace_failures=parser.trace_numeric_failures, enforce_mappings=True,
        )
    enforce_quality(diagnostics, quality_policy)
    tables = tuple(prepare_table(snapshot) for snapshot in parser.table_snapshots)
    messages = list(diagnostics.warnings)
    if not tables:
        messages.append('No tables were emitted by the document parser.')
    if isinstance(source, bytes):
        hash_input, hash_kind = source, 'original bytes'
    elif source_url is None:
        hash_input, hash_kind = source.encode('utf-8'), 'supplied text UTF-8'
    else:
        hash_input, hash_kind = html.encode('utf-8'), 'normalized HTML UTF-8'
    payload, worksheets = render_workbook(
        tables, source_url=link_url, source_hash=hashlib.sha256(hash_input).hexdigest(),
        hash_kind=hash_kind, document_diagnostics=messages,
    )
    results = tuple(XlsxTableResult(w.ordinal, w.worksheet_name, w.status, w.issues)
                    for w in worksheets)
    issues = tuple(f'Table {table.ordinal}: {issue}' for table in results for issue in table.issues)
    if quality_policy == 'strict' and issues:
        raise XlsxQualityError(issues)
    messages.extend(issues)
    status = ('no_tables' if not results else 'needs_review'
              if messages or any(t.status != 'exported' for t in results) else 'complete')
    _publish(payload, path, overwrite=overwrite, load_workbook=load_workbook)
    return XlsxExportResult(path, status, results, tuple(messages))
