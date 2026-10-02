"""Read-only: find percent-formatted cells in RCQ workbooks whose row label is not a percent measure."""
import glob, os, re, sys, warnings
warnings.filterwarnings("ignore")
from openpyxl import load_workbook
PCT_LABEL = re.compile(r"%|percent|margin|rate|growth|change|yield|ratio|share of", re.I)
for path in sorted(glob.glob("E:/RCQWealth/*/Originals/SEC/*/*.xlsx")):
    name = os.path.basename(path)
    if name.startswith("~$"): continue
    wb = load_workbook(path, read_only=True, data_only=True)
    suspicious, pct_cells, sheets_hit, samples = 0, 0, set(), []
    for ws in wb.worksheets[1:]:
        for row in ws.iter_rows():
            label = next((str(c.value) for c in row if isinstance(c.value, str) and c.value.strip()), "")
            for c in row:
                if isinstance(c.value, (int, float)) and "%" in (c.number_format or ""):
                    pct_cells += 1
                    if not PCT_LABEL.search(label) and abs(c.value) >= 2:
                        suspicious += 1; sheets_hit.add(ws.title)
                        if len(samples) < 3: samples.append(f"{ws.title}!{c.coordinate} '{label[:40]}' = {c.value} [{c.number_format}]")
    wb.close()
    print(f"{name[:40]:40} sheets={len(wb.sheetnames)-1:3} pct_cells={pct_cells:4} suspicious={suspicious:4} in {len(sheets_hit)} sheets")
    for s in samples: print("      ", s)
