"""Read-only verification of the created RDDT workbooks and repeat-run selection."""
import hashlib
import importlib.util
import json
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
report = json.loads((HERE / "folder-export.json").read_text(encoding="utf-8"))
checks = []
for item in report["results"]:
    if item["status"] != "created":
        continue
    source = Path(item["source"])
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    assert digest == item["sha256"], source
    with Path(item["workbook"]).open("rb") as stream:
        workbook = load_workbook(stream)
        metadata = dict(workbook["Contents"].iter_rows(
            min_row=1, max_row=6, max_col=2, values_only=True))
        assert metadata["SHA-256"] == digest
        assert metadata["SHA-256 kind"] == "original bytes"
        assert len(workbook.worksheets) == item["tables"] + 1
        assert workbook["Contents"].freeze_panes is None
        assert all(sheet.freeze_panes == "B1" for sheet in workbook.worksheets[1:])
        checks.append({"workbook": item["workbook"], "source_sha256": digest,
                       "worksheets": len(workbook.worksheets), "verified": True})
        workbook.close()

helper_path = Path(r"C:\Users\einstein\.codex\skills\xbrl-to-excel\scripts\export_folder.py")
spec = importlib.util.spec_from_file_location("export_folder", helper_path)
helper = importlib.util.module_from_spec(spec)
spec.loader.exec_module(helper)

def unexpected_export(*args, **kwargs):
    raise AssertionError("Repeat run must not invoke export")

rows = helper.run(Path(r"E:\RCQWealth\RDDT\Originals\SEC\2024"), unexpected_export,
                  dry_run=True)
assert sum(row["status"] == "skipped_identical" for row in rows) == 3
assert not any(row["status"] in {"error", "would_create", "skipped_existing"} for row in rows)
print(json.dumps({"workbooks": checks, "repeat_run_skipped_identical": 3}, indent=2))
