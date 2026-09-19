from dataclasses import replace
from io import BytesIO

import pytest
from bs4 import BeautifulSoup
from openpyxl import load_workbook

from sec2md.xlsx_tables import prepare_table, snapshot_html_table
from sec2md.xlsx_values import CellValue
from sec2md.xlsx_writer import render_workbook


def sample():
    source = '''<h2>Income statement</h2><p>Amounts in millions</p><table id="native">
    <tr><th></th><th colspan="2">Three Months Ended</th></tr>
    <tr><th>Item</th><th>Jul 26, 2026</th><th>Jul 27, 2025</th></tr>
    <tr><td>Revenue</td><td>96,221</td><td>46,743</td></tr></table>'''
    soup = BeautifulSoup(source, 'lxml')
    return prepare_table(snapshot_html_table(soup.table, ordinal=1, page=3, source_url=None))


def render(tables, **kwargs):
    payload, records = render_workbook(tables, source_url='https://example.com/filing.htm',
                                      source_hash='abc123', hash_kind='utf8_text', **kwargs)
    return load_workbook(BytesIO(payload), data_only=False), records


def values(sheet):
    return [c.value for row in sheet for c in row if c.value is not None]


def test_grouped_headers_saved_values_navigation_and_no_frozen_rows():
    table = sample()
    wb, records = render([table])
    assert len(wb.worksheets) == 2
    record, = records
    sheet = wb[record.worksheet_name]
    assert record.ordinal == 1 and record.status == table.status
    assert wb.worksheets[0].title == 'Contents'
    assert wb.worksheets[0].freeze_panes is None
    assert sheet.freeze_panes == 'B1'
    assert 96221 in values(sheet)
    assert '96,221' in values(sheet)
    for heading in table.headers[1:]:
        assert heading in values(sheet)
    assert 'Amounts in millions' in values(sheet)
    assert any(c.hyperlink and c.hyperlink.target == 'https://example.com/filing.htm#native'
               for row in sheet for c in row)
    links = [c.hyperlink.location for row in wb['Contents'] for c in row if c.hyperlink]
    assert len(links) == 2
    assert all(record.worksheet_name in link for link in links)
    for ws in wb:
        assert not ws.auto_filter.ref
        assert not any(c.data_type == 'f' for row in ws for c in row)
    assert not sheet.merged_cells.ranges


@pytest.mark.parametrize('text', ['=1+1', '+SUM(A1)', '-1+2', '@SUM(A1)', '#N/A'])
def test_source_strings_are_explicit_text_everywhere(text):
    table = sample()
    table = replace(table, title=text, units=text, headers=(text,),
                    rows=((CellValue(text, '@', text),),), notes=(text,),
                    source=replace(table.source, original_text=text))
    wb, records = render([table])
    cells = [c for ws in wb for row in ws for c in row if c.value == text]
    assert len(cells) >= 5
    assert all(c.data_type == 's' for c in cells)


def test_names_order_fallback_and_native_anchor_only():
    table = sample()
    tables = [replace(table, title="'Contents[]:*?/\\" * 4,
                      source=replace(table.source, ordinal=n, source_anchor=None,
                                     element_id='sec2md-generated'),
                      status='source_text_only', issues=('Bad geometry',), rows=())
              for n in (2, 1, 1)]
    wb, records = render(tables)
    names = [r.worksheet_name for r in records]
    assert [r.ordinal for r in records] == [2, 1, 1]
    assert len(set(n.lower() for n in names)) == 3
    assert all(len(n) <= 31 and not any(c in n for c in '[]:*?/\\') for n in names)
    for name in names:
        sheet = wb[name]
        assert any('Copy grid unavailable' in str(v) for v in values(sheet))
        assert 'Bad geometry' in values(sheet)
        assert not any(c.hyperlink and 'sec2md-generated' in str(c.hyperlink.target)
                       for row in sheet for c in row)


def test_long_text_chunks_reconstruct_and_report_review():
    text = '=start' + 'abcdef ' * 10000
    table = sample()
    cell = replace(table.source.source_cells[0], text=text)
    table = replace(table, source=replace(table.source, source_cells=(cell,), original_text=text),
                    rows=((CellValue(text, '@', text),),), headers=('Text',))
    wb, records = render([table])
    sheet = wb[records[0].worksheet_name]
    assert records[0].issues and records[0].status == 'needs_review'
    assert all(len(v) <= 32767 for v in values(sheet) if isinstance(v, str))
    # Original extracted text is a labeled, consecutive sequence of exact chunks.
    start = next(c.row for row in sheet for c in row if c.value == 'Original extracted text (chunks; concatenate without separators)')
    chunks = []
    for row in sheet.iter_rows(min_row=start + 1):
        if row[0].value is None:
            break
        chunks.append(row[0].value)
    assert ''.join(chunks) == text


def test_empty_workbook_has_honest_hash_and_document_diagnostics():
    wb, records = render([], document_diagnostics=('Parse warning',))
    assert not records and wb.sheetnames == ['Contents']
    assert 'utf8_text' in values(wb['Contents'])
    assert 'abc123' in values(wb['Contents'])
    assert 'Parse warning' in values(wb['Contents'])
    assert any('No tables' in str(v) for v in values(wb['Contents']))


def test_dimension_limit_retains_text_and_reports_writer_issue(monkeypatch):
    import sec2md.xlsx_writer as writer
    monkeypatch.setattr(writer, 'MAX_COLUMNS', 2)
    wb, records = render([sample()])
    assert records[0].status == 'source_text_only'
    assert any('dimension' in issue.lower() for issue in records[0].issues)
    assert any('96,221' in str(v) for v in values(wb[records[0].worksheet_name]))


def test_original_grid_and_span_ledger_are_visible():
    table = sample()
    wb, records = render([table])
    sheet = wb[records[0].worksheet_name]
    assert any('Original cell grid' in str(v) for v in values(sheet))
    numeric_original = next(c for row in sheet for c in row if c.value == '96,221')
    assert numeric_original.column == 2
    assert numeric_original.data_type == 's'
    assert any('rowspan 1, colspan 2' in str(v) for v in values(sheet))


def test_formats_blank_dash_zero_and_original_provenance():
    from decimal import Decimal
    table = replace(sample(), headers=('Label', 'Amount', 'Percent', 'Blank', 'Dash'),
                    rows=((CellValue('Example', '@', 'Example'),
                           CellValue(Decimal('0'), '#,##0', '0'),
                           CellValue(Decimal('.123'), '0.0%', '12.3%'),
                           CellValue(None, '@', ''), CellValue('—', '@', '—')),))
    wb, records = render([table])
    sheet = wb[records[0].worksheet_name]
    row = next(row for row in sheet if row[0].value == 'Example')
    assert [c.value for c in row[:5]] == ['Example', 0, .123, None, '—']
    assert row[2].number_format == '0.0%'
    assert any('Body row 0, column 1:' in str(v) for v in values(sheet))


def test_xml_incompatible_text_is_reversibly_escaped_and_reported():
    table = replace(sample(), title='Bad\x01 title')
    wb, records = render([table])
    assert records[0].status == 'needs_review'
    assert any('XML' in issue for issue in records[0].issues)
    assert any('\\u0001' in str(v) for v in values(wb[records[0].worksheet_name]))


def test_row_limit_degrades_to_explicit_unavailable_in_full(monkeypatch):
    import sec2md.xlsx_writer as writer
    monkeypatch.setattr(writer, 'MAX_ROWS', 40)
    table = replace(sample(), notes=('Long note ' * 2000,))
    wb, records = render([table])
    sheet = wb[records[0].worksheet_name]
    assert sheet.max_row <= 40
    assert any('UNAVAILABLE IN FULL' in str(v) for v in values(sheet))
    assert records[0].status == 'source_text_only'
    assert any('unavailable in full' in issue.lower() for issue in records[0].issues)


def test_table_has_original_navigation_and_normalization_disclosure():
    wb, records = render([sample()])
    sheet = wb[records[0].worksheet_name]
    assert any(c.hyperlink and c.hyperlink.location and records[0].worksheet_name in c.hyperlink.location
               for row in sheet for c in row)
    assert any('whitespace' in str(v).lower() for v in values(sheet))


def test_writer_import_does_not_load_optional_dependency():
    import subprocess
    import sys
    result = subprocess.run([sys.executable, '-c',
                             'import sys; import sec2md.xlsx_writer; assert "openpyxl" not in sys.modules'],
                            capture_output=True, text=True)
    assert result.returncode == 0, result.stderr


def test_supplied_title_and_printed_page_are_retained():
    table = replace(sample(), title='Explicit full filing title',
                    source=replace(sample().source, display_page=17))
    wb, records = render([table])
    sheet = wb[records[0].worksheet_name]
    assert sheet['A1'].value == table.title
    assert 'Parser page 3 | Printed page 17' in values(sheet)
    assert 17 in values(wb['Contents'])
