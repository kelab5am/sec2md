"""Struck-through text renders as GFM strikethrough when it holds a letter or digit.

Blacklined credit agreements show an amendment as a struck deletion directly followed by
an underlined insertion, with no whitespace between them. Joined as the source writes it,
Salesforce's "84</font><font>86" read "8486" and "LIBOR" + "the Term SOFR" read
"LIBORthe Term SOFR". The ``~~`` markers are the boundary: "~~84~~86". Struck text with no
letter or digit (a full stop, a bullet, a dash) is a formatting accident and renders as
before, joined as the source writes it.
"""

from decimal import Decimal

import pytest
from bs4 import BeautifulSoup

from sec2md import convert_to_markdown
from sec2md.parser import Parser
from sec2md.table_parser import TableParser, extract_cell_text, inline_text, render_cell_content
from sec2md.xlsx_values import convert_cell

TNR = "font-family:'Times New Roman';font-size:12pt"
STRUCK = f'<font style="color:#ff0000;{TNR};text-decoration:line-through">'
INSERTED = f'<font style="color:#0000ff;{TNR};text-decoration:underline">'
PLAIN = f'<font style="{TNR}">'
URL = "https://www.sec.gov/Archives/edgar/data/1/000000000100000001/repro.htm"

# repro-blackline.htm, from Salesforce's EX-10.9 (0001108524-22-000022).
REPRO_BLACKLINE = (
    "<html><body>"
    f"<div>{PLAIN}Section 7.07 Judgments </font>{STRUCK}84</font>{INSERTED}86</font></div>"
    f"<div>{PLAIN}with respect to any amounts denominated in Dollars, </font>{STRUCK}LIBOR</font>"
    f"{INSERTED}the Term SOFR Reference Rate</font>"
    f"{PLAIN}, (ii) with respect to any amounts denominated in Sterling</font></div>"
    f'<div>{PLAIN}Split capital: </font>{PLAIN}D</font><font style="font-size:9pt">ELOITTE</font></div>'
    "</body></html>"
)

# repro-strike-other.htm: struck text with no letter or digit, and Tesla's struck phrase.
REPRO_STRIKE_OTHER = (
    "<html><body>"
    '<div><span style="font-family:Arial;font-size:10pt">as reflected by the credit derivative'
    " notional value. See Note 3 for further details on credit derivatives</span>"
    '<span style="font-family:Arial;font-size:10pt;text-decoration:line-through">.</span>'
    '<span style="font-family:Arial;font-size:10pt"> The maximum exposure</span></div>'
    '<div><font style="font-family:Arial;font-size:10pt;font-weight:700;text-decoration:line-through">'
    "&#8226;</font>"
    '<font style="font-family:Arial;font-size:10pt;font-weight:700;padding-left:14.5pt">'
    "FY25 revenue of $37.9 billion</font></div>"
    '<p>criteria established in <span style="font-style:italic"><span style="text-decoration:'
    'line-through;">Internal Control - Integrated Framework</span></span> (2013) issued by COSO.</p>'
    "<table><tr><td><font>In default</font></td><td><font style=\"text-decoration:line-through\">-"
    "</font></td><td><font>1,234</font></td></tr></table>"
    "</body></html>"
)


def _convert(html: str) -> str:
    pages = convert_to_markdown(
        html, base_url=URL, return_pages=True, embed_images=False, quality_policy="warn"
    )
    return "\n\n".join(page.content for page in pages)


def _content(body: str) -> str:
    parser = Parser(f"<html><body>{body}</body></html>")
    return "\n\n".join(page.content for page in parser.get_pages(include_elements=False))


def _cell(html: str):
    return BeautifulSoup(f"<table><tr>{html}</tr></table>", "lxml").td


class TestRepros:
    def test_blackline_repro(self):
        assert _convert(REPRO_BLACKLINE) == (
            "Section 7.07 Judgments ~~84~~86\n\n"
            "with respect to any amounts denominated in Dollars, ~~LIBOR~~the Term SOFR"
            " Reference Rate, (ii) with respect to any amounts denominated in Sterling\n\n"
            "Split capital: DELOITTE"
        )

    def test_strike_other_repro(self):
        # Unchanged from a09a4fb, but for Tesla's struck phrase, which is now marked.
        assert _convert(REPRO_STRIKE_OTHER) == (
            "as reflected by the credit derivative notional value. See Note 3 for further"
            " details on credit derivatives. The maximum exposure\n\n"
            "**• FY25 revenue of $37.9 billion**\n\n"
            "criteria established in *~~Internal Control - Integrated Framework~~* (2013)"
            " issued by COSO.\n\n"
            "In default - 1,234"
        )


class TestStream:
    @pytest.mark.parametrize("tag", ["s", "strike", "del"])
    def test_strike_tags_mark_like_line_through(self, tag):
        assert _content(f"<p>clauses (a) through (<{tag}>w</{tag}>x) above</p>") == (
            "clauses (a) through (~~w~~x) above"
        )

    @pytest.mark.parametrize(
        "style",
        [
            "text-decoration:line-through",
            "text-decoration: line-through",
            "TEXT-DECORATION:LINE-THROUGH",
            "text-decoration-line:line-through",
            "text-decoration:underline line-through",
        ],
    )
    def test_line_through_style_variants(self, style):
        assert _content(f'<p>Section <span style="{style}">84</span><u>86</u></p>') == (
            "Section ~~84~~86"
        )

    def test_underline_alone_has_no_marker(self):
        body = '<p>the <span style="text-decoration:underline">Term SOFR</span> rate</p>'
        assert _content(body) == "the Term SOFR rate"

    def test_source_whitespace_stays_outside_the_markers(self):
        assert _content("<p>rate<s> LIBOR </s>plus</p>") == "rate ~~LIBOR~~ plus"

    def test_adjacent_struck_spans_merge_into_one_run(self):
        assert _content("<p>the <s>Euro</s><s>currency</s>SOFR Rate</p>") == (
            "the ~~Eurocurrency~~SOFR Rate"
        )

    def test_nested_strikes_mark_once(self):
        body = '<p>the <s><span style="text-decoration:line-through">LIBOR</span> Rate</s>SOFR</p>'
        assert _content(body) == "the ~~LIBOR Rate~~SOFR"

    def test_struck_and_bold_on_one_element(self):
        body = '<p><span style="font-weight:700;text-decoration:line-through">84</span>86</p>'
        assert _content(body) == "**~~84~~**86"

    @pytest.mark.parametrize("text", [".", "-", "&#8226;", "&#8212;", " . ", "()"])
    def test_struck_text_without_letter_or_digit_is_unchanged(self, text):
        assert _content(f"<p>derivatives<s>{text}</s> The</p>") == _content(
            f"<p>derivatives<span>{text}</span> The</p>"
        )

    def test_struck_text_in_a_positioned_line(self):
        body = (
            '<div style="position:relative;width:600px;height:800px">'
            '<div style="position:absolute;left:10px;top:10px;font-size:10pt">'
            "<span>Judgments </span><s>84</s><span>86</span></div></div>"
        )
        assert "Judgments ~~84~~86" in _content(body)

    def test_struck_positioned_line(self):
        body = (
            '<div style="position:relative;width:600px;height:800px">'
            '<div style="position:absolute;left:10px;top:10px;font-size:10pt;'
            'text-decoration:line-through">EXECUTION VERSION</div>'
            '<div style="position:absolute;left:10px;top:40px;font-size:10pt">EXHIBIT A</div></div>'
        )
        assert "~~EXECUTION VERSION~~" in _content(body)


class TestEmphasis:
    def test_struck_then_inserted_inside_bold(self):
        assert _content(f"<p><b>Judgments {STRUCK}84</font>{INSERTED}86</font></b></p>") == (
            "**Judgments ~~84~~86**"
        )

    def test_struck_then_inserted_inside_italic(self):
        assert _content(f"<p><i>in Dollars, {STRUCK}LIBOR</font>{INSERTED}SOFR</font></i></p>") == (
            "*in Dollars, ~~LIBOR~~SOFR*"
        )

    def test_struck_inside_italic_as_tesla_writes_it(self):
        body = (
            '<p>in <span style="font-style:italic"><span style="text-decoration:line-through;">'
            "Internal Control - Integrated Framework</span></span> (2013)</p>"
        )
        assert _content(body) == "in *~~Internal Control - Integrated Framework~~* (2013)"

    def test_struck_bold_runs_meeting_merge(self):
        assert _content("<p><b><s>Euro</s></b><b><s>currency</s></b>SOFR</p>") == (
            "**~~Eurocurrency~~**SOFR"
        )

    def test_struck_bullet_in_bold_is_unchanged(self):
        body = (
            '<div><font style="font-weight:700;text-decoration:line-through">&#8226;</font>'
            '<font style="font-weight:700;padding-left:14.5pt">FY25 revenue</font></div>'
        )
        assert _content(body) == "**• FY25 revenue**"


class TestTableCells:
    def test_inline_text_marks_the_struck_span(self):
        assert inline_text(_cell(f"<td>Judgments {STRUCK}84</font>{INSERTED}86</font></td>")) == (
            "Judgments ~~84~~86"
        )

    def test_extract_cell_text_marks_the_struck_span(self):
        cell = _cell(f"<td>{STRUCK}LIBOR</font>{INSERTED}the Term SOFR</font></td>")
        assert extract_cell_text(cell) == "~~LIBOR~~the Term SOFR"

    def test_whitespace_in_a_struck_cell_span_stays_outside(self):
        assert inline_text(_cell("<td>rate<s> LIBOR </s>plus</td>")) == "rate ~~LIBOR~~ plus"

    def test_adjacent_struck_spans_in_a_cell_merge(self):
        assert inline_text(_cell("<td><s>Euro</s><s>currency</s>SOFR</td>")) == "~~Eurocurrency~~SOFR"

    def test_struck_cell_itself(self):
        assert inline_text(_cell('<td style="text-decoration:line-through">84</td>')) == "~~84~~"

    def test_linked_struck_text_keeps_the_markers_in_the_label(self):
        cell = _cell('<td><a href="#s7"><s>84</s>86</a></td>')
        assert render_cell_content(cell) == "[~~84~~86](#s7)"

    def test_struck_text_around_a_link(self):
        cell = _cell('<td><s><a href="#s7">84</a></s>86</td>')
        assert render_cell_content(cell) == "~~[84](#s7)~~86"

    def test_struck_dash_in_a_cell_is_unchanged(self):
        assert inline_text(_cell("<td><s>-</s></td>")) == "-"

    def test_table_render_marks_the_struck_span(self):
        table = BeautifulSoup(
            "<table>"
            "<tr><td>Section</td><td>Page</td></tr>"
            f"<tr><td>Judgments</td><td>{STRUCK}84</font>{INSERTED}86</font></td></tr>"
            "</table>",
            "lxml",
        ).table
        assert "~~84~~86" in TableParser(table).md()

    def test_document_table_traces_struck_numbers(self):
        parser = Parser(
            "<html><body><table>"
            "<tr><td>Section</td><td>Page</td></tr>"
            f"<tr><td>Judgments</td><td>{STRUCK}84</font>{INSERTED}86</font></td></tr>"
            "<tr><td>Total</td><td><s>1,234</s></td></tr>"
            "</table></body></html>"
        )
        content = parser.get_pages(include_elements=True)[0].content
        assert "~~84~~86" in content and "~~1,234~~" in content
        assert parser.trace_numeric_failures == ()


class TestStruckNumberValues:
    """A struck number in a cell is deleted text, not a value: XLSX keeps it as text.

    The XLSX snapshot reads cells from the HTML (xlsx_tables._visible_text), not from the
    Markdown, so it marks struck text the same way: "84" struck before "86" would otherwise
    export as the number 8486.
    """

    def test_struck_number_is_not_converted(self):
        cell = convert_cell("~~1,234~~", role="number")
        assert cell.value == "~~1,234~~"
        assert cell.number_format == "@"
        assert cell.review_reason is not None

    def test_inserted_number_still_converts(self):
        assert convert_cell("1,234", role="number").value == Decimal("1234")

    @pytest.mark.parametrize(
        "amount,text",
        [
            ("<s>1,234</s>", "~~1,234~~"),
            (f"{STRUCK}84</font>{INSERTED}86</font>", "~~84~~86"),
            ("<s>1,2</s><s>34</s>", "~~1,234~~"),
        ],
    )
    def test_xlsx_keeps_a_struck_number_as_reviewed_text(self, tmp_path, amount, text):
        from openpyxl import load_workbook

        from sec2md import export_xlsx
        from sec2md.xlsx_tables import prepare_table, snapshot_html_table

        html = (
            "<table><tr><th>Item</th><th>Amount</th></tr>"
            f"<tr><td>Revenue</td><td>{amount}</td></tr>"
            "<tr><td>Cost</td><td>567</td></tr></table>"
        )
        snapshot = snapshot_html_table(
            BeautifulSoup(html, "html.parser").table, ordinal=1, page=1, source_url=None
        )
        assert snapshot.source_cells[3].text == text
        prepared = prepare_table(snapshot)
        assert prepared.rows[0][1].value == text
        assert prepared.rows[0][1].review_reason is not None
        assert prepared.rows[1][1].value == Decimal("567")
        result = export_xlsx(html, tmp_path / "struck.xlsx")
        assert result.status == "needs_review"
        workbook = load_workbook(result.path)
        assert text in [cell.value for row in workbook.worksheets[1].iter_rows() for cell in row]
        workbook.close()

    def test_xlsx_struck_dash_is_unchanged(self):
        from sec2md.xlsx_tables import snapshot_html_table

        html = "<table><tr><td>In default</td><td><s>-</s></td></tr><tr><td>A</td><td>1</td></tr></table>"
        snapshot = snapshot_html_table(
            BeautifulSoup(html, "html.parser").table, ordinal=1, page=1, source_url=None
        )
        assert snapshot.source_cells[1].text == "-"
