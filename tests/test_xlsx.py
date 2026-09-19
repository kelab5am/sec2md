"""Public XLSX orchestration and publication, using real saved workbooks."""
import builtins
from dataclasses import FrozenInstanceError
import hashlib
import os
import subprocess
import sys

import pytest
from openpyxl import load_workbook

import sec2md
from sec2md.utils import FetchedHtml


TABLE = ('<table id="sales"><tr><th>Item</th><th>2026</th></tr>'
         '<tr><td>Revenue</td><td>96,221</td></tr></table>')
BAD = ('<table><tr><td colspan="bad">Unreliable</td><td>123</td></tr>'
       '<tr><td>Tail</td><td>456</td></tr></table>')


def values(workbook):
    return [cell.value for sheet in workbook for row in sheet for cell in row
            if cell.value is not None]


def test_public_api_is_available():
    assert hasattr(sec2md, 'export_xlsx')
    for name in ('XlsxExportResult', 'XlsxTableResult', 'XlsxQualityError',
                 'XlsxDependencyError'):
        assert name in sec2md.__all__


@pytest.mark.parametrize('as_bytes', [True, False])
def test_offline_round_trip_and_hash(tmp_path, monkeypatch, as_bytes):
    monkeypatch.setattr('sec2md.core.fetch', lambda *a, **k: pytest.fail('offline'))
    source = TABLE.encode() if as_bytes else TABLE
    result = sec2md.export_xlsx(source, tmp_path / 'filing.xlsx')
    assert result.path == tmp_path / 'filing.xlsx'
    assert result.status == 'complete'
    assert len(result.tables) == 1
    assert result.tables[0].ordinal == 1
    assert result.tables[0].status == 'exported'
    wb = load_workbook(result.path)
    assert wb.sheetnames == ['Contents', result.tables[0].sheet_name]
    assert 96221 in values(wb)
    assert '96,221' in values(wb)
    assert any(hashlib.sha256(TABLE.encode()).hexdigest() in str(v) for v in values(wb))
    assert any(('original bytes' if as_bytes else 'supplied text UTF-8') in str(v)
               for v in values(wb))
    assert wb.worksheets[1].freeze_panes == 'B1'
    wb.close()
    assert list(tmp_path.iterdir()) == [result.path]


def test_base_url_resolves_without_fetch(tmp_path, monkeypatch):
    monkeypatch.setattr('sec2md.core.fetch', lambda *a, **k: pytest.fail('offline'))
    html = TABLE.replace('Revenue', '<a href="note.htm">Revenue</a>') + '<img src="x.png">'
    result = sec2md.export_xlsx(html, tmp_path / 'a.xlsx', base_url='https://example.com/a.htm')
    wb = load_workbook(result.path)
    assert any('https://example.com/note.htm' in str(v) for v in values(wb))
    wb.close()


def test_dependency_checked_before_fetch_and_write(tmp_path, monkeypatch):
    original = builtins.__import__
    def blocked(name, *args, **kwargs):
        if name == 'openpyxl':
            raise ImportError('absent')
        return original(name, *args, **kwargs)
    monkeypatch.setattr(builtins, '__import__', blocked)
    monkeypatch.setattr('sec2md.core.fetch', lambda *a, **k: pytest.fail('fetch'))
    with pytest.raises(sec2md.XlsxDependencyError, match=r'sec2md\[xlsx\]'):
        sec2md.export_xlsx('https://example.com/a.htm', tmp_path / 'a.xlsx')
    assert not list(tmp_path.iterdir())


def test_package_import_without_extra():
    code = '''
import builtins
old = builtins.__import__
def blocked(name, *a, **k):
    if name.startswith('openpyxl'):
        raise ImportError('blocked')
    return old(name, *a, **k)
builtins.__import__ = blocked
from sec2md import export_xlsx, XlsxExportResult, XlsxTableResult
'''
    subprocess.run([sys.executable, '-c', code], check=True)


@pytest.mark.parametrize('policy', ['warn', 'off'])
def test_partial_recovery_keeps_every_table(tmp_path, policy):
    result = sec2md.export_xlsx(BAD + TABLE, tmp_path / 'a.xlsx', quality_policy=policy)
    assert result.status == 'needs_review'
    assert [t.ordinal for t in result.tables] == [1, 2]
    assert result.tables[0].issues
    assert result.tables[1].status == 'exported'
    assert result.diagnostics
    wb = load_workbook(result.path)
    assert len(wb.sheetnames) == 3
    assert any('Unreliable' in str(v) for v in values(wb))
    wb.close()


def test_strict_rejects_table_issues_before_publication(tmp_path):
    target = tmp_path / 'a.xlsx'
    target.write_bytes(b'old')
    with pytest.raises(sec2md.XlsxQualityError) as caught:
        sec2md.export_xlsx(BAD + TABLE, target, quality_policy='strict', overwrite=True)
    assert caught.value.issues
    assert target.read_bytes() == b'old'
    assert list(tmp_path.iterdir()) == [target]


def test_no_tables_is_diagnostic_workbook(tmp_path):
    result = sec2md.export_xlsx('<p>Ordinary prose</p>', tmp_path / 'a.xlsx')
    assert result.status == 'no_tables'
    assert not result.tables
    assert result.diagnostics
    wb = load_workbook(result.path)
    assert wb.sheetnames == ['Contents']
    assert any('No tables' in str(v) for v in values(wb))
    wb.close()


def test_existing_destination_is_preserved(tmp_path):
    target = tmp_path / 'a.xlsx'
    target.write_bytes(b'old')
    with pytest.raises(FileExistsError):
        sec2md.export_xlsx(TABLE, target)
    assert target.read_bytes() == b'old'
    assert list(tmp_path.iterdir()) == [target]


def test_concurrent_destination_never_replaced(tmp_path, monkeypatch):
    target = tmp_path / 'a.xlsx'
    link = os.link
    def race(src, dst):
        target.write_bytes(b'concurrent')
        link(src, dst)
    monkeypatch.setattr('sec2md.xlsx.os.link', race)
    with pytest.raises(FileExistsError):
        sec2md.export_xlsx(TABLE, target)
    assert target.read_bytes() == b'concurrent'
    assert list(tmp_path.iterdir()) == [target]


@pytest.mark.parametrize('operation', ['link', 'replace'])
def test_publication_failure_cleans_temp_and_preserves_target(tmp_path, monkeypatch, operation):
    target = tmp_path / 'a.xlsx'
    if operation == 'replace':
        target.write_bytes(b'old')
    def fail(*a, **k):
        raise PermissionError('locked or unsupported filesystem')
    monkeypatch.setattr(f'sec2md.xlsx.os.{operation}', fail)
    with pytest.raises(OSError, match='locked or unsupported'):
        sec2md.export_xlsx(TABLE, target, overwrite=operation == 'replace')
    assert list(tmp_path.iterdir()) == ([target] if operation == 'replace' else [])
    if operation == 'replace':
        assert target.read_bytes() == b'old'


def test_overwrite_saved_workbook(tmp_path):
    target = tmp_path / 'a.xlsx'
    target.write_bytes(b'old')
    sec2md.export_xlsx(TABLE, target, overwrite=True)
    wb = load_workbook(target)
    assert 96221 in values(wb)
    wb.close()


@pytest.mark.parametrize('case', ['missing_parent', 'directory', 'invalid_policy', 'base_url'])
def test_invalid_arguments_fail_before_fetch(tmp_path, monkeypatch, case):
    monkeypatch.setattr('sec2md.core.fetch', lambda *a, **k: pytest.fail('fetch'))
    target = tmp_path / 'a.xlsx'
    kwargs = {}
    error = ValueError
    if case == 'missing_parent':
        target = tmp_path / 'missing' / 'a.xlsx'
        error = FileNotFoundError
    elif case == 'directory':
        target = tmp_path
        error = IsADirectoryError
    elif case == 'invalid_policy':
        kwargs['quality_policy'] = 'invalid'
    else:
        kwargs['base_url'] = 'http://example.com/a.htm'
    with pytest.raises(error):
        sec2md.export_xlsx('https://example.com/a.htm', target, **kwargs)


def test_url_fetches_only_primary_and_labels_normalized_hash(tmp_path, monkeypatch):
    calls = []
    def fetch(url, user_agent=None):
        calls.append((url, user_agent))
        return FetchedHtml(content=(TABLE + '<img src="image.png">').encode(), charset=None)
    monkeypatch.setattr('sec2md.core.fetch', fetch)
    result = sec2md.export_xlsx('https://example.com/a.htm', tmp_path / 'a.xlsx', user_agent='test')
    assert calls == [('https://example.com/a.htm', 'test')]
    wb = load_workbook(result.path)
    assert any('normalized HTML UTF-8' in str(v) for v in values(wb))
    wb.close()


def test_strict_honors_writer_issues(tmp_path):
    source = TABLE.replace('Revenue', 'R' * 401)
    target = tmp_path / 'a.xlsx'
    with pytest.raises(sec2md.XlsxQualityError) as caught:
        sec2md.export_xlsx(source, target, quality_policy='strict')
    assert any('chunk' in issue.lower() for issue in caught.value.issues)
    assert not list(tmp_path.iterdir())
    result = sec2md.export_xlsx(source, target, quality_policy='off')
    assert result.status == 'needs_review'
    assert result.tables[0].issues


def test_strict_propagates_parse_quality_error(tmp_path):
    target = tmp_path / 'a.xlsx'
    target.write_bytes(b'old')
    with pytest.raises(sec2md.ParseQualityError):
        sec2md.export_xlsx('<p>Damaged \ufffd source</p>' + TABLE, target,
                           quality_policy='strict', overwrite=True)
    assert target.read_bytes() == b'old'
    assert list(tmp_path.iterdir()) == [target]


@pytest.mark.parametrize('failure', ['render', 'parse', 'reopen', 'write'])
def test_prepublication_failures_preserve_existing_file(tmp_path, monkeypatch, failure):
    target = tmp_path / 'a.xlsx'
    target.write_bytes(b'old')
    def fail(*args, **kwargs):
        raise RuntimeError('unexpected fault')
    if failure == 'render':
        monkeypatch.setattr('sec2md.xlsx.render_workbook', fail)
    elif failure == 'parse':
        monkeypatch.setattr('sec2md.xlsx.Parser.get_pages', fail)
    elif failure == 'reopen':
        monkeypatch.setattr('openpyxl.load_workbook', fail)
    else:
        import tempfile
        create = tempfile.NamedTemporaryFile
        class FailingFile:
            def __init__(self, file):
                self.file = file
                self.name = file.name
            def __enter__(self):
                return self
            def __exit__(self, *args):
                self.file.close()
            def write(self, content):
                self.file.write(content[:10])
                raise RuntimeError('unexpected fault')
        monkeypatch.setattr('sec2md.xlsx.tempfile.NamedTemporaryFile',
                            lambda **kwargs: FailingFile(create(**kwargs)))
    with pytest.raises(RuntimeError, match='unexpected fault'):
        sec2md.export_xlsx(TABLE, target, overwrite=True)
    assert target.read_bytes() == b'old'
    assert list(tmp_path.iterdir()) == [target]


def test_corrupt_render_is_not_published(tmp_path, monkeypatch):
    from zipfile import BadZipFile
    import sec2md.xlsx as api
    render = api.render_workbook
    def corrupt(*args, **kwargs):
        _, mapping = render(*args, **kwargs)
        return b'not a workbook', mapping
    monkeypatch.setattr(api, 'render_workbook', corrupt)
    with pytest.raises(BadZipFile):
        sec2md.export_xlsx(TABLE, tmp_path / 'a.xlsx')
    assert not list(tmp_path.iterdir())


def test_parser_runs_once_with_images_disabled(tmp_path, monkeypatch):
    import sec2md.xlsx as api
    parse = api.Parser.get_pages
    calls = []
    def track(self, *args, **kwargs):
        calls.append(kwargs)
        return parse(self, *args, **kwargs)
    monkeypatch.setattr(api.Parser, 'get_pages', track)
    sec2md.export_xlsx(TABLE + '<img src="image.png">', tmp_path / 'a.xlsx')
    assert calls == [{'include_images': False}]


def test_results_are_frozen_and_blank_dash_zero_are_reliable(tmp_path):
    html = (TABLE.replace('96,221', '0').replace('</table>',
            '<tr><td>Blank</td><td></td></tr><tr><td>Dash</td><td>—</td></tr></table>'))
    result = sec2md.export_xlsx(html, tmp_path / 'a.xlsx', quality_policy='strict')
    assert result.status == 'complete'
    assert result.tables[0].issues == ()
    with pytest.raises(FrozenInstanceError):
        result.status = 'needs_review'
    with pytest.raises(FrozenInstanceError):
        result.tables[0].status = 'needs_review'
    wb = load_workbook(result.path)
    assert 0 in values(wb)
    assert '—' in values(wb)
    wb.close()


def test_original_bytes_hash_precedes_decoding(tmp_path):
    source = TABLE.replace('Revenue', 'Caf\u00e9').encode('cp1252')
    result = sec2md.export_xlsx(source, tmp_path / 'a.xlsx')
    wb = load_workbook(result.path)
    assert any(hashlib.sha256(source).hexdigest() in str(v) for v in values(wb))
    assert 'Café' in values(wb)
    wb.close()
