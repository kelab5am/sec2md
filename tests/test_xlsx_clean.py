from io import BytesIO

from bs4 import BeautifulSoup
from openpyxl import load_workbook
import pytest

from sec2md import export_xlsx
from sec2md.xlsx_tables import prepare_table, snapshot_html_table
from sec2md.xlsx_writer import render_workbook


TABLE = '<table><tr><th>Item</th><th>2025</th></tr><tr><td>Revenue</td><td>123</td></tr></table>'


def test_styled_filing_title_skips_units_and_unaudited():
    soup = BeautifulSoup('<div><span style="font-weight:700">Consolidated Balance Sheets</span></div>'
                         '<div>(in thousands)</div><div>(unaudited)</div>' + TABLE, 'lxml')
    snapshot = snapshot_html_table(soup.table, ordinal=1, page=1, source_url=None)
    assert snapshot.explicit_title == 'Consolidated Balance Sheets'


def test_table_intro_is_a_source_grounded_fallback_title():
    intro = 'The following table represents our revenue disaggregated by geography:'
    soup = BeautifulSoup(f'<p>{intro}</p>{TABLE}', 'lxml')
    snapshot = snapshot_html_table(soup.table, ordinal=1, page=1, source_url=None)
    assert snapshot.explicit_title == 'Revenue disaggregated by geography'


def test_explanatory_footnote_becomes_short_table_title():
    intro = '(1)Adjusted cost of revenue is cost of revenue adjusted for taxes as follows:'
    soup = BeautifulSoup(f'<p>{intro}</p>{TABLE}', 'lxml')
    snapshot = snapshot_html_table(soup.table, ordinal=1, page=1, source_url=None)
    assert snapshot.explicit_title == 'Adjusted cost of revenue'


def test_italic_title_is_preferred_over_bold_units():
    soup = BeautifulSoup('<div><i>Average revenue per user</i></div>'
                         '<div><b>(in dollars)</b></div>' + TABLE, 'lxml')
    snapshot = snapshot_html_table(soup.table, ordinal=1, page=1, source_url=None)
    assert snapshot.explicit_title == 'Average revenue per user'


def test_clean_workbook_has_tables_without_diagnostic_sections():
    soup = BeautifulSoup('<h2>Income statement</h2><p>In millions</p>' + TABLE, 'lxml')
    table = prepare_table(snapshot_html_table(soup.table, ordinal=1, page=1, source_url=None))
    payload, records = render_workbook([table], source_url=None, source_hash='abc', hash_kind='original bytes')
    wb = load_workbook(BytesIO(payload))
    ws = wb[records[0].worksheet_name]
    assert ws.title == 'Income statement'
    assert all(s.freeze_panes is None for s in wb)
    visible = [str(c.value) for s in wb for row in s if not s.row_dimensions[row[0].row].hidden
               for c in row if c.value is not None]
    for unwanted in ('Notes and review', 'source cells', 'extracted text', 'source origins',
                     'Status:', 'Nearby source context', 'Local element:', 'SHA-256', 'Copy A'):
        assert not any(unwanted in v for v in visible)
    assert 123 in [c.value for row in ws for c in row]
    assert '123' in [c.value for row in ws for c in row]


def test_front_matter_ends_at_actual_contents_table_not_page_number(tmp_path):
    toc = '<h2>Table of Contents</h2><table>' + ''.join(
        f'<tr><td><a href="#s{i}">Item {i}. Financial statements</a></td><td>{i}</td></tr>'
        for i in range(1, 4)) + '</table>'
    html = TABLE + toc + '<h2>Income statement</h2>' + TABLE
    result = export_xlsx(html, tmp_path / 'clean.xlsx')
    assert len(result.tables) == 1
    assert result.tables[0].ordinal == 3


def test_repeated_contents_navigation_does_not_remove_financial_tables(tmp_path):
    html = '<a href="#toc">Table of Contents</a><h2>Income statement</h2>' + TABLE
    result = export_xlsx(html, tmp_path / 'keep.xlsx')
    assert len(result.tables) == 1


def test_small_filing_without_contents_keeps_early_tables(tmp_path):
    result = export_xlsx(TABLE, tmp_path / 'early.xlsx')
    assert len(result.tables) == 1


def test_untitled_table_uses_descriptive_source_row_label():
    soup = BeautifulSoup(TABLE, 'lxml')
    snapshot = snapshot_html_table(soup.table, ordinal=7, page=1, source_url=None)
    assert prepare_table(snapshot).title == 'Revenue'


@pytest.mark.parametrize('inside', [True, False])
def test_contents_heading_inside_table_or_with_empty_anchor(tmp_path, inside):
    title = ('<tr><td colspan="2">Table of Contents</td></tr>' if inside else '')
    preceding = '' if inside else '<div><a id="toc"></a><b>Table of Contents</b></div>'
    toc = preceding + '<table>' + title + ''.join(
        f'<tr><td><a href="#s{i}">Item {i}. Financial statements</a></td><td>{i}</td></tr>'
        for i in range(1, 4)) + '</table>'
    result = export_xlsx(TABLE + toc + TABLE, tmp_path / 'toc.xlsx')
    assert [t.ordinal for t in result.tables] == [3]


def test_unsupported_geometry_rejects_without_publishing_incomplete_workbook(tmp_path):
    html = '<table><tr><td rowspan="0">Cash</td><td>123</td></tr><tr><td>456</td></tr></table>'
    target = tmp_path / 'invalid.xlsx'
    with pytest.raises(ValueError, match='grid unavailable'):
        export_xlsx(html, target)
    assert not target.exists()
